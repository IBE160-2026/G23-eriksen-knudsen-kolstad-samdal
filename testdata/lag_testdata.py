"""Lager Excel-mal, testeksempel og fasit for COSI - MRP.

Bruk: uv run --with openpyxl python lag_testdata.py <testdata-mappe> [--alle]

Uten --alle lages bare malen. Med --alle lages også testeksempel 01 og fasit.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

UT = Path(sys.argv[1])
MAKS_UKER = 12

FONT = Font(name="Arial", size=10)
FONT_FET = Font(name="Arial", size=10, bold=True)
FONT_TITTEL = Font(name="Arial", size=14, bold=True)
FONT_HVIT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FYLL_HODE = PatternFill("solid", fgColor="1F4E78")
FYLL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FYLL_EKSEMPEL = PatternFill("solid", fgColor="E2EFDA")
FYLL_VARSEL = PatternFill("solid", fgColor="F8CBAD")
FYLL_KRITISK = PatternFill("solid", fgColor="FCE4D6")
TYNN = Side(style="thin", color="BFBFBF")
RAMME = Border(left=TYNN, right=TYNN, top=TYNN, bottom=TYNN)

STYKK_KOLONNER = ["Nivå", "Varenummer", "Beskrivelse", "Antall", "Type",
                  "Lagerbeholdning", "Sikkerhetslager", "Ledetid (uker)"]
NIVAER = ["0", ".1", "..2"]
TYPER = ["Ferdigvare", "Komponent (Kjøpes inn)", "Komponent (Produseres selv)"]
KJOPT = "Komponent (Kjøpes inn)"
PROD = "Komponent (Produseres selv)"

# Eksempelet i Veiledning: skruesettet brukes både i setet og i ryggen, og lager mm. står bare første gang
EKSEMPEL_STYKKLISTE = [
    ("0", "100-001", "Kontorstol Basic", None, "Ferdigvare", 10, 0, 1),
    (".1", "200-001", "Sete", 1, PROD, 40, 10, 1),
    ("..2", "300-001", "Seteplate", 1, KJOPT, 50, 10, 2),
    ("..2", "300-002", "Skruesett M6", 4, KJOPT, 150, 40, 2),
    (".1", "200-002", "Rygg", 1, PROD, 30, 5, 1),
    ("..2", "300-002", "Skruesett M6", 2, KJOPT, None, None, None),
    (".1", "300-003", "Hjul", 5, KJOPT, 80, 30, 1),
]


# ---------- Testeksempel 01: fiktiv kontorstolprodusent ----------

PRODUKT = "Kontorstol Basic"
MPS = [20, 25, 30, 20, 35, 30, 25, 40, 30, 20]  # 10 ukers horisont

# (varenummer, beskrivelse, antall per ferdigprodukt, lager, sikkerhetslager, ledetid, hensikt)
# Testeksempel 01 er enkeltnivå: alle komponentene er innkjøpt og sitter på nivå .1.
KOMPONENTER = [
    ("300-001", "Sete", 1, 200, 20, 2, "Normalforløp: første behov kan bestilles i tide"),
    ("300-002", "Rygg", 1, 60, 10, 1, "Normalforløp med ledetid 1 uke"),
    ("300-003", "Hjul", 5, 80, 30, 1, "Bestilling forsinket i uke 1 (lager går under 0)"),
    ("300-004", "Armlene", 2, 50, 20, 1, "Bruk av sikkerhetslager i uke 1 (lager mellom 0 og sikkerhetslager)"),
    ("300-005", "Gassfjær", 1, 75, 10, 2, "Mottak forsinket: behov i uke 3, planlagt mottak først i uke 4"),
    ("300-006", "Stolfot", 1, 400, 20, 10, "Ledetid = horisont (grensetilfelle)"),
    ("300-007", "Skruesett", 4, 150, 40, 2, "Planlagt mottak i uke 2 kommer i tide (ingen eksepsjon)"),
    ("300-008", "Trekk", 1, 90, 0, 3, "Sikkerhetslager 0 og frigivelse nøyaktig i uke 1 (grensetilfelle)"),
]

# (varenummer, uke, antall)
MOTTAK = [("300-005", 4, 100), ("300-007", 2, 120)]


def beregn(mps, komp, mottak):
    """MRP for én komponent etter reglene i arket 'Regler'."""
    _id, _navn, per, lager, ss, ledetid = komp[:6]
    h = len(mps)
    brutto = [m * per for m in mps]
    planlagt = [0] * h
    for kid, uke, antall in mottak:
        if kid == _id:
            planlagt[uke - 1] += antall
    forventet, netto, ordremottak, frigivelse = [], [], [], [0] * h
    eksepsjoner = []
    forfalte = []  # (indeks i eksepsjoner, forfalt antall, frigi_uke)
    if ledetid >= h:
        eksepsjoner.append((_id, "–", "Ledetid overskrider horisont",
                            f"Ledetid {ledetid} uker ≥ horisont {h} uker. "
                            "Ingen ny bestilling kan dekke behov innenfor horisonten."))
    forrige = lager
    for t in range(h):
        uke = t + 1
        disponibel = forrige + planlagt[t] - brutto[t]
        nb = max(0, ss - disponibel)
        netto.append(nb)
        ordremottak.append(nb)
        if nb > 0:
            frigi_uke = uke - ledetid
            if frigi_uke >= 1:
                frigivelse[frigi_uke - 1] += nb
            else:
                frigivelse[0] += nb  # forfalt: må frigis straks
                # Meldingen trenger total frigivelse i uke 1, som først er kjent etter løkken
                forfalte.append((len(eksepsjoner), nb, frigi_uke))
                if disponibel < 0:
                    eksepsjoner.append((_id, uke, "Bestilling forsinket",
                                        f"Lager blir {disponibel} i uke {uke}. Bestillingen på {{nb}} måtte vært "
                                        f"frigitt i uke {frigi_uke}, før horisonten. {{forfalt}} "
                                        "Leveranse til kunde blir forsinket."))
                else:
                    eksepsjoner.append((_id, uke, "Bruk av sikkerhetslager",
                                        f"Lager blir {disponibel} i uke {uke}, under sikkerhetslager {ss}, "
                                        f"før ny leveranse kan komme inn. Bestillingen på {{nb}} måtte vært "
                                        f"frigitt i uke {frigi_uke}. {{forfalt}}"))
            senere = [u for u in range(uke + 1, h + 1) if planlagt[u - 1] > 0]
            if senere:
                eksepsjoner.append((_id, uke, "Mottak forsinket",
                                    f"Nettobehov {nb} i uke {uke}, mens planlagt mottak ligger i uke {senere[0]}. "
                                    f"Vurder å fremskynde mottaket."))
        forrige = disponibel + nb
        forventet.append(forrige)
    for i, antall, _ in forfalte:
        kid, uke, typ, melding = eksepsjoner[i]
        forfalt = (f"De {antall} er forfalt og tatt med i frigivelsen på {frigivelse[0]} i uke 1, "
                   f"men kommer tidligst i uke {1 + ledetid}.")
        eksepsjoner[i] = (kid, uke, typ, melding.format(nb=antall, forfalt=forfalt))
    rader = {
        "Bruttobehov": brutto,
        "Planlagte mottak": planlagt,
        "Forventet lagerbeholdning": forventet,
        "Nettobehov": netto,
        "Planlagte ordremottak": ordremottak,
        "Planlagte ordrefrigivelser": frigivelse,
    }
    return rader, eksepsjoner


# ---------- Hjelpefunksjoner for formatering ----------

def hode(ws, rad, verdier, kol=1):
    for i, v in enumerate(verdier):
        c = ws.cell(rad, kol + i, v)
        c.font, c.fill, c.border = FONT_HVIT, FYLL_HODE, RAMME
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def celle(ws, rad, kol, verdi=None, fyll=None, fet=False):
    c = ws.cell(rad, kol, verdi)
    c.font = FONT_FET if fet else FONT
    c.border = RAMME
    if fyll:
        c.fill = fyll
    return c


def bredder(ws, bredder_):
    for bokstav, b in bredder_.items():
        ws.column_dimensions[bokstav].width = b


def heltall_validering(ws, omraade, minimum=0):
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=str(minimum), allow_blank=True,
                        showErrorMessage=True, errorTitle="Ugyldig verdi",
                        error=f"Skriv et heltall større enn eller lik {minimum}.")
    ws.add_data_validation(dv)
    dv.add(omraade)


def tekstblokk(ws, start, linjer, kol=1):
    r = start
    for linje in linjer:
        c = ws.cell(r, kol, linje)
        c.font = FONT_FET if linje.endswith(":") else FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return r


# ---------- Malen (inndatafilen brukeren laster opp) ----------

PROGNOSE = "MPS"
PROGNOSE_REF = f"'{PROGNOSE}'"
PROGNOSE_START = 8    # uker som vises før brukeren har fylt ut noe
MOTTAK_RADER = 200
STYKK_RADER = 1000    # én rad per del; samme del kan stå flere steder i stykklisten
FONT_LENKE = Font(name="Arial", size=10, color="0563C1", underline="single")
FONT_SKJULT = Font(name="Arial", size=10, color="FFFFFF")
FONT_KOMMER = Font(name="Arial", size=10, italic=True, color="808080")
FONT_ROD = Font(name="Arial", size=10, color="C00000")
# Fyll i betinget formatering må ha både start- og sluttfarge for å vises riktig i Excel
AKTIV_RAD = dict(fill=PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
                 border=RAMME, font=Font(color="000000"))
UTFYLT = dict(fill=PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF"), border=RAMME)
GRAA = dict(fill=PatternFill("solid", start_color="E7E6E6", end_color="E7E6E6"), border=RAMME)


def lenke(ws, ref, tekst, ark, celle_="A1"):
    c = ws[ref]
    c.value = tekst
    c.hyperlink = Hyperlink(ref=ref, location=f"'{ark}'!{celle_}", display=tekst)
    c.font = FONT_LENKE
    c.alignment = Alignment(vertical="top")
    return c


def tilbake_lenke(ws, ref):
    lenke(ws, ref, "← Til veiledning", "Veiledning")


def hvit_naar_utfylt(ws, omraade):
    """Gule celler blir hvite når de er fylt ut. Må legges til før vis_rad_naar, så den vinner."""
    topp = omraade.split(":")[0]
    ws.conditional_formatting.add(omraade, FormulaRule(formula=[f"NOT(ISBLANK({topp}))"],
                                                       stopIfTrue=True, **UTFYLT))


def vis_rad_naar(ws, omraade, formel, gul=True):
    """Rader i området blir gule med ramme når formelen er sann, ellers ser de blanke ut."""
    stil = AKTIV_RAD if gul else {k: v for k, v in AKTIV_RAD.items() if k != "fill"}
    ws.conditional_formatting.add(omraade, FormulaRule(formula=[formel], **stil))


def graa_naar(ws, omraade, formel):
    """Celler som skal stå tomme blir grå. Må legges til før vis_rad_naar, så den vinner."""
    ws.conditional_formatting.add(omraade, FormulaRule(formula=[formel], stopIfTrue=True, **GRAA))


def regel_validering(ws, omraade, formel, feil):
    dv = DataValidation(type="custom", formula1=formel, allow_blank=True,
                        showErrorMessage=True, errorTitle="Ugyldig verdi", error=feil)
    ws.add_data_validation(dv)
    dv.add(omraade)


def tekstrader(ws, start, rader, sluttkol=7):
    """Skriver rader med (lenke/etikett i kolonne A, tekst i B–sluttkol).

    Etiketten er enten None, en tekst eller (tekst, arknavn) for lenke til arket.
    """
    r = start
    for etikett, tekst in rader:
        if isinstance(etikett, tuple):
            lenke(ws, f"A{r}", *etikett)
        elif etikett:
            c = ws.cell(r, 1, etikett)
            c.font = FONT_FET
            c.alignment = Alignment(vertical="top")
        c = ws.cell(r, 2, tekst)
        c.font = FONT_KOMMER if str(tekst).startswith("Kommer:") else FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=sluttkol)
        ws.row_dimensions[r].height = 15 * (len(str(tekst or "")) // 85 + 1)
        r += 1
    return r


def overskrift(ws, rad, tekst):
    ws.cell(rad, 1, tekst).font = FONT_FET
    return rad + 1


def kontroll(n):
    """Formel for Kontroll-kolonnen i Stykkliste: hva som mangler eller ikke stemmer i rad n."""
    p, nx = n - 1, n + 1
    forste = f"COUNTIF($B$1:$B{p},$B{n})=0"
    typer = "{" + ",".join(f'"{t}"' for t in TYPER) + "}"
    if n == 2:
        nivaregel = [f'IF(AND($A{n}<>"",$A{n}&""<>"0"),"Første rad må være nivå 0. ","")']
    else:
        nivaregel = [
            f'IF($A{n}&""="0","Bare første rad kan være nivå 0. ","")',
            f'IF(AND($A{n}&""="..2",$A{p}&""<>".1",$A{p}&""<>"..2"),"Nivå ..2 må stå rett under .1 eller ..2. ","")',
        ]
    deler = [
        f'IF($A{n}="","Mangler nivå. ","")',
        # Innlimte verdier omgår nedtrekkslistene, så verdiene kontrolleres også her
        f'IF(AND($A{n}<>"",ISNA(MATCH($A{n}&"",$Z$2:$Z$4,0))),"Ugyldig nivå, velg 0, .1 eller ..2. ","")',
        *nivaregel,
        f'IF($B{n}="","Mangler varenummer. ","")',
        f'IF($C{n}="","Mangler beskrivelse. ","")',
        f'IF(AND($A{n}<>"",$A{n}&""<>"0",$D{n}=""),"Mangler antall. ","")',
        f'IF($E{n}="","Mangler type. ","")',
        f'IF(AND($E{n}<>"",ISNA(MATCH($E{n},{typer},0))),"Ugyldig type, velg fra listen. ","")',
        f'IF(AND($B{n}<>"",{forste},COUNT($F{n}:$H{n})<3),"Mangler lager, sikkerhetslager eller ledetid. ","")',
        f'IF(AND($B{n}<>"",NOT({forste}),COUNTIFS($B$1:$B{p},$B{n},$C$1:$C{p},$C{n})=0),'
        f'"Annen beskrivelse enn første gang varenummeret står. ","")',
        f'IF(AND($B{n}<>"",NOT({forste}),COUNTIFS($B$1:$B{p},$B{n},$E$1:$E{p},$E{n})=0),'
        f'"Annen type enn første gang varenummeret står. ","")',
        # Nivåene er 0, .1 og ..2, så lengre tekst betyr lavere nivå
        f'IF(AND(ISNUMBER(SEARCH("Kjøpes inn",$E{n})),$A{nx}<>"",LEN($A{nx})>LEN($A{n})),'
        f'"Har underkomponenter, men kjøpes inn. ","")',
        f'IF(AND($A{n}&""="0",$E{n}<>"",$E{n}<>"Ferdigvare"),"Nivå 0 skal ha type Ferdigvare. ","")',
        f'IF(AND($A{n}<>"",$A{n}&""<>"0",$E{n}="Ferdigvare"),"Ferdigvare brukes bare på nivå 0. ","")',
    ]
    return f'=IF(COUNTA($A{n}:$H{n})=0,"",' + "&".join(deler) + ")"


def lag_mal(mps=None, stykkliste=None, mottak=None):
    wb = Workbook()
    eksempel = mps is not None

    # Veiledning
    ws = wb.active
    ws.title = "Veiledning"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "COSI - MRP: Excel-mal for inndata"
    ws["A1"].font = FONT_TITTEL
    r = overskrift(ws, 3, "Slik fyller du ut malen (klikk på arknavnet for å gå dit):")
    r = tekstrader(ws, r, [
        ("Generelt", "Fyll kun ut gule celler. En celle blir hvit når den er fylt ut, og grå celler skal stå tomme. "
                     "Ikke endre arknavn, kolonneoverskrifter eller rekkefølge. Nye rader blir gule etter hvert som "
                     "du fyller ut."),
        (("Stykkliste", "Stykkliste"), "Ferdigproduktet og alle delene det består av, med lagerstatus, "
                                       "sikkerhetslager og ledetid. Én rad per del, ovenfra og ned. Nivå 0 er "
                                       "ferdigproduktet (første rad, ett produkt per fil). Nivå .1 er "
                                       "hovedkomponentene, f.eks. sete, rygg og armlene. Nivå ..2 er delene som "
                                       "trengs for å bygge hovedkomponenten rett over. Neste rad kommer frem når "
                                       "nivå og varenummer i raden over er fylt ut."),
        ((PROGNOSE, PROGNOSE), "Antall ferdigprodukter per uke. Fyll ut fra uke 1 og sammenhengende fremover. "
                               f"De {PROGNOSE_START} første ukene vises fra start, og neste uke kommer frem når "
                               f"alle ukene over er fylt ut. Horisonten er antall utfylte uker (8–{MAKS_UKER})."),
        (("Planlagte mottak", "Planlagte mottak"), "Bestillinger som allerede er lagt hos leverandør, og "
                                                   "produksjonsordrer som allerede er startet. Én rad per "
                                                   "mottak. Varenummer velges fra en liste som henter "
                                                   "varenumrene fra Stykkliste og oppdateres av seg selv. "
                                                   "Neste rad kommer frem når raden over er fylt ut. Arket "
                                                   "kan stå tomt hvis det ikke finnes planlagte mottak."),
        (("Veien videre", "Veien videre"), "Hva du gjør når malen er ferdig utfylt."),
    ], sluttkol=8)
    r = overskrift(ws, r + 1, "Regler for nivåer i stykklisten:")
    r = tekstrader(ws, r, [
        (None, "Velg nivå fra nedtrekkslisten. Første rad må være 0, og det kan bare være ett nivå 0."),
        (None, "Nivå .1 kan bare brukes når nivå 0 er lagt inn. Nivå ..2 kan bare brukes rett under en .1 "
               "eller en annen ..2."),
        (None, "Du kan gå opp igjen fra ..2 til .1 når delene til en hovedkomponent er lagt inn. En del på nivå "
               "..2 tilhører nærmeste .1 over seg."),
        (None, "Kolonnen «Kontroll» sier fra hvis en rad bryter reglene over."),
    ], sluttkol=8)
    r = overskrift(ws, r + 1, "Regler for nedtrekkslister:")
    r = tekstrader(ws, r, [
        (None, "Ikke lim inn data i celler med nedtrekksliste: Nivå og Type i Stykkliste og Varenummer i "
               "Planlagte mottak. Innliming fjerner listen og kontrollen av verdien i cellen. Velg verdien fra "
               "listen eller skriv den inn. Har du limt inn ved et uhell, angrer du med Ctrl+Z."),
    ], sluttkol=8)
    r = overskrift(ws, r + 1, "Regler for verdier:")
    r = tekstrader(ws, r, [
        (("Stykkliste", "Stykkliste"), "Samme del kan stå flere steder i stykklisten, f.eks. skruer. Den står "
                                       "da med samme varenummer, beskrivelse og type hver gang."),
        (("Stykkliste", "Stykkliste"), "Antall er hvor mange av delen som trengs per enhet av delen over, "
                                       "heltall ≥ 1. Antall fylles ikke ut for nivå 0."),
        (("Stykkliste", "Stykkliste"), "Type velges fra listen: " + ", ".join(TYPER) + ". Ferdigvare brukes "
                                       "bare på nivå 0. En del med underkomponenter kan ikke være en type som "
                                       "kjøpes inn."),
        (("Stykkliste", "Stykkliste"), "Lagerbeholdning (ved starten av uke 1) og sikkerhetslager er heltall ≥ 0. "
                                       "Ledetid er hele uker: innkjøpstid for deler som kjøpes inn og "
                                       "produksjonstid for deler som produseres selv, heltall ≥ 1 (0 er lov for "
                                       "ferdigvaren). De tre fylles "
                                       "bare ut første gang et varenummer står i stykklisten."),
        (("Stykkliste", "Stykkliste"), "Kolonnen «Kontroll» viser hva som mangler eller ikke stemmer i hver rad, "
                                       "bl.a. brudd på reglene for nivåer og verdier som ikke finnes i "
                                       "nedtrekkslistene. Den skal være tom når stykklisten er ferdig."),
        ((PROGNOSE, PROGNOSE), "Antall ferdigprodukter er heltall ≥ 0."),
        (("Planlagte mottak", "Planlagte mottak"), "Varenummer velges fra listen. Uke må ligge innenfor "
                                                   "horisonten (1 til siste utfylte uke)."),
    ], sluttkol=8)
    r = overskrift(ws, r + 1, "Forslag til varenummer:")
    r = tekstrader(ws, r, [
        (None, "Tre siffer for varegruppe, bindestrek og tre siffer løpenummer, f.eks. 300-002. 100 = "
               "ferdigvarer, 200 = komponenter som produseres selv, 300 = komponenter som kjøpes inn. Nummeret sier ikke noe om nivå, "
               "slik at samme del kan brukes flere steder og flyttes i stykklisten uten å få nytt nummer. "
               "Bindestreken gjør at Excel lagrer nummeret som tekst."),
    ], sluttkol=8)
    r = overskrift(ws, r + 1, "Eksempel (grønn = kun eksempel, skal ikke kopieres inn). "
                              "Skruesettet brukes både i setet og i ryggen:")
    hode(ws, r, STYKK_KOLONNER)
    ws.row_dimensions[r].height = 30
    for j, rad in enumerate(EKSEMPEL_STYKKLISTE):
        for i, v in enumerate(rad):
            celle(ws, r + 1 + j, i + 1, v, FYLL_EKSEMPEL)
    bredder(ws, {"A": 22, "B": 14, "C": 18, "D": 10, "E": 26, "F": 16, "G": 16, "H": 12})

    # Stykkliste: én rad per del ovenfra og ned. Neste rad vises når nivå og varenummer i raden over er
    # fylt ut (overskriftsraden teller som utfylt, så rad 2 vises fra start)
    ws = wb.create_sheet("Stykkliste")
    ws.sheet_view.showGridLines = False
    hode(ws, 1, STYKK_KOLONNER)
    siste = STYKK_RADER + 1
    for i in range(STYKK_RADER):
        n = i + 2
        rad = stykkliste[i] if eksempel and i < len(stykkliste) else [None] * len(STYKK_KOLONNER)
        for k, v in enumerate(rad):
            ws.cell(n, k + 1, v).font = FONT
        ws.cell(n, 1).number_format = "@"  # ellers gjør Excel .1 om til tallet 0,1
        ws.cell(n, 10, kontroll(n)).font = FONT_ROD
        # Varenummerlisten til Planlagte mottak (skjulte kolonner): AA nummererer første forekomst av hvert
        # varenummer, AB henter varenumrene i den rekkefølgen
        ws.cell(n, 27, f'=IF(AND($B{n}<>"",COUNTIF($B$2:$B{n},$B{n})=1),MAX($AA$1:$AA{n - 1})+1,"")')
        ws.cell(n, 28, f'=IFERROR(INDEX($B$2:$B${siste},MATCH(ROW()-1,$AA$2:$AA${siste},0)),"")')
    forrige_fylt = "COUNTA($A1:$B1)=2"
    forste = "COUNTIF($B$1:$B1,$B2)=0"
    hvit_naar_utfylt(ws, f"A2:H{siste}")
    graa_naar(ws, f"D2:D{siste}", '$A2&""="0"')
    graa_naar(ws, f"F2:H{siste}", f'AND($B2<>"",NOT({forste}))')
    vis_rad_naar(ws, f"A2:H{siste}", forrige_fylt)
    # Nivålisten ligger i en skjult kolonne. Rad 2 kan bare være 0, raden rett under 0 bare .1
    for n, niva in enumerate(NIVAER, start=2):
        ws.cell(n, 26, niva).number_format = "@"
    for kol in ["Z", "AA", "AB"]:
        ws.column_dimensions[kol].hidden = True
    # Navnet Varenumre er så langt som antall ulike varenumre, så listen vokser når stykklisten fylles ut
    wb.defined_names["Varenumre"] = DefinedName(
        "Varenumre", attr_text=f"OFFSET(Stykkliste!$AB$2,0,0,MAX(1,MAX(Stykkliste!$AA$2:$AA${siste})),1)")
    dv = DataValidation(type="list", formula1='IF(ROW()=2,$Z$2,IF($A1&""="0",$Z$3,$Z$3:$Z$4))',
                        allow_blank=True, showErrorMessage=True, errorTitle="Ugyldig nivå",
                        error="Velg nivå fra listen. Første rad er nivå 0 (ferdigproduktet). Nivå .1 krever at "
                              "nivå 0 er lagt inn, og ..2 må stå rett under en .1 eller en annen ..2.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{siste}")
    regel_validering(ws, f"B2:B{siste}", f'AND({forrige_fylt},$A2<>"")',
                     "Fyll ut nivå og varenummer i raden over først, og velg nivå i denne raden.")
    regel_validering(ws, f"C2:C{siste}", forrige_fylt, "Fyll ut nivå og varenummer i raden over først.")
    regel_validering(ws, f"D2:D{siste}", f'AND({forrige_fylt},$A2&""<>"0",D2=INT(D2),D2>=1)',
                     "Antall fylles ikke ut for nivå 0 (ferdigproduktet). Ellers må antall være et heltall "
                     "større enn eller lik 1.")
    dv = DataValidation(type="list", formula1='"' + ",".join(TYPER) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Ugyldig type", error="Velg type fra listen.")
    ws.add_data_validation(dv)
    dv.add(f"E2:E{siste}")
    regel_validering(ws, f"F2:G{siste}", f"AND({forrige_fylt},{forste},F2=INT(F2),F2>=0)",
                     "Fylles bare ut første gang varenummeret står i stykklisten. Verdien må være et heltall "
                     "større enn eller lik 0.")
    regel_validering(ws, f"H2:H{siste}", f'AND({forrige_fylt},{forste},H2=INT(H2),H2>=IF($A2&""="0",0,1))',
                     "Fylles bare ut første gang varenummeret står i stykklisten. Ledetid må være et heltall "
                     "større enn eller lik 1 (0 er lov for ferdigproduktet).")
    bredder(ws, {"A": 8, "B": 13, "C": 24, "D": 9, "E": 26, "F": 15, "G": 15, "H": 11, "I": 4, "J": 50,
                 "K": 4, "L": 18})
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws["J1"] = "Kontroll"
    ws["J1"].font = FONT_FET
    ws["J1"].alignment = Alignment(vertical="center")
    tilbake_lenke(ws, "L1")

    # MPS
    ws = wb.create_sheet(PROGNOSE)
    ws.sheet_view.showGridLines = False
    hode(ws, 1, ["Uke", "Antall ferdigprodukter"])
    for u in range(1, MAKS_UKER + 1):
        verdi = mps[u - 1] if eksempel and u <= len(mps) else None
        if u <= PROGNOSE_START:
            celle(ws, u + 1, 1, u, fet=True)
            celle(ws, u + 1, 2, verdi, FYLL_INPUT)
        else:
            ws.cell(u + 1, 1, u).font = FONT_SKJULT
            ws.cell(u + 1, 2, verdi).font = FONT
        ws.cell(u + 1, 1).alignment = Alignment(horizontal="center")
    forste_skjulte = PROGNOSE_START + 2
    # Uke u vises når alle ukene over er fylt ut
    vis = f"COUNT($B$1:$B{forste_skjulte - 1})=ROW()-2"
    hvit_naar_utfylt(ws, f"B2:B{MAKS_UKER + 1}")
    vis_rad_naar(ws, f"A{forste_skjulte}:A{MAKS_UKER + 1}", vis, gul=False)
    vis_rad_naar(ws, f"B{forste_skjulte}:B{MAKS_UKER + 1}", vis)
    heltall_validering(ws, f"B2:B{forste_skjulte - 1}")
    regel_validering(ws, f"B{forste_skjulte}:B{MAKS_UKER + 1}",
                     f"AND(COUNT($B$1:$B{forste_skjulte - 1})=ROW()-2,"
                     f"B{forste_skjulte}=INT(B{forste_skjulte}),B{forste_skjulte}>=0)",
                     "Fyll ut alle ukene over først. Verdien må være et heltall større enn eller lik 0.")
    bredder(ws, {"A": 8, "B": 24, "C": 4, "D": 18})
    ws.freeze_panes = "A2"
    tilbake_lenke(ws, "D1")

    # Planlagte mottak: neste rad vises når raden over er fylt ut (overskriftsraden teller som utfylt)
    ws = wb.create_sheet("Planlagte mottak")
    ws.sheet_view.showGridLines = False
    hode(ws, 1, ["Varenummer", "Uke", "Antall"])
    siste = MOTTAK_RADER + 1
    for i in range(MOTTAK_RADER):
        rad = mottak[i] if eksempel and i < len(mottak) else [None] * 3
        for k, v in enumerate(rad):
            if i == 0:
                celle(ws, 2, k + 1, v, FYLL_INPUT)
            else:
                ws.cell(i + 2, k + 1, v).font = FONT
    hvit_naar_utfylt(ws, f"A2:C{siste}")
    vis_rad_naar(ws, f"A3:C{siste}", "COUNTA($A2:$C2)=3")
    forrige_fylt = "COUNTA($A1:$C1)=3"
    horisont = f"COUNT({PROGNOSE_REF}!$B$2:$B${MAKS_UKER + 1})"
    dv = DataValidation(type="list", formula1="Varenumre", allow_blank=True, showErrorMessage=True,
                        errorTitle="Ugyldig varenummer", error="Velg et varenummer fra listen. Listen viser "
                                                               "varenumrene i arket Stykkliste.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{siste}")
    regel_validering(ws, f"B2:B{siste}", f"AND({forrige_fylt},B2=INT(B2),B2>=1,B2<={horisont})",
                     "Fyll ut raden over først. Uke må være et heltall fra 1 til siste utfylte uke "
                     f"i {PROGNOSE}.")
    regel_validering(ws, f"C2:C{siste}", f"AND({forrige_fylt},C2=INT(C2),C2>=1)",
                     "Fyll ut raden over først. Antall må være et heltall større enn eller lik 1.")
    bredder(ws, {"A": 14, "B": 8, "C": 10, "D": 4, "E": 18})
    ws.freeze_panes = "A2"
    tilbake_lenke(ws, "E1")

    # Veien videre: skjelett som fylles ut etter hvert som appen utvikles
    ws = wb.create_sheet("Veien videre")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Veien videre: når malen er ferdig utfylt"
    ws["A1"].font = FONT_TITTEL
    tilbake_lenke(ws, "A2")
    c = ws.cell(3, 1, "Utkast. Denne veiledningen bygges ut etter hvert som appen utvikles. "
                      "Grå tekst merket «Kommer:» er ikke bestemt ennå.")
    c.font = FONT_KOMMER
    r = overskrift(ws, 5, "Steg for steg:")
    r = tekstrader(ws, r, [
        ("1. Kontroller", "Se over at alle gule celler er fylt ut, at kolonnen Kontroll i Stykkliste er tom, og at "
                          "ukene i MPS er sammenhengende fra uke 1."),
        ("2. Lagre", "Lagre filen som .xlsx. Filnavnet kan du velge selv, men arknavnene må ikke endres."),
        ("3. Last opp", "Last opp filen i COSI - MRP."),
        (None, "Kommer: lenke til appen og hvordan opplastingen gjøres."),
        ("4. Rett feil", "Appen kontrollerer at nødvendige felt og verdier er fylt ut før beregningen starter. "
                         "Rett eventuelle feil i malen og last opp på nytt."),
        (None, "Kommer: oversikt over feilmeldinger og hvordan de rettes."),
        ("5. Les resultatet", "Resultatet vises i en tabell per uke og komponent: bruttobehov, planlagte mottak, "
                              "forventet lagerbeholdning, nettobehov, planlagte ordremottak og planlagte "
                              "ordrefrigivelser. Komponenter med minst ett varsel er markert som kritiske."),
        (None, "Kommer: skisse av resultattabellen og forklaring av fremhevingene."),
        ("6. Håndter varsler", "Hvert varsel gjelder en uke og en komponent (eller hele komponenten):"),
        (None, "• Bruk av sikkerhetslager: lageret går under sikkerhetslageret før ny leveranse kan komme inn."),
        (None, "• Bestilling forsinket: behovet kan ikke dekkes i tide, og leveransen blir forsinket (til kunden "
                "for ferdigproduktet, eller til produksjonen av delen over i stykklisten)."),
        (None, "• Mottak forsinket: et planlagt mottak kommer senere enn behovet og bør fremskyndes."),
        (None, "• Ledetid overskrider horisont: ledetiden er like lang som eller lengre enn horisonten."),
        (None, "Kommer: anbefalte tiltak for hvert varsel."),
        ("7. Ved endringer", "Endres noe i planen, oppdaterer du malen og laster den opp på nytt. Appen viser da "
                             "hva som er nytt, endret eller fjernet sammenlignet med forrige kjøring."),
        (None, "Kommer: hvordan før/etter-sammenligningen leses."),
    ])
    bredder(ws, {"A": 22, "B": 14, "C": 16, "D": 16, "E": 16, "F": 16, "G": 10})
    return wb


# ---------- Fasit ----------

REGLER = [
    "Beregningsregler (bekreftet av gruppen):",
    "Notasjon: uke t = 1..H, lager før uke 1 = lagerbeholdning, SS = sikkerhetslager, L = ledetid.",
    "1. Bruttobehov(t) = MPS(t) × antall per ferdigprodukt.",
    "2. Disponibel(t) = forventet lager(t−1) + planlagte mottak(t) − bruttobehov(t).",
    "3. Nettobehov(t) = maks(0, SS − disponibel(t)).",
    "4. Lot-for-lot: planlagt ordremottak(t) = nettobehov(t).",
    "5. Forventet lagerbeholdning(t) = disponibel(t) + planlagt ordremottak(t).",
    "6. Planlagt ordrefrigivelse i uke t − L. Hvis t − L < 1 er frigivelsen forfalt og legges i uke 1 (må frigis straks), "
    "sammen med eventuelle ordinære frigivelser for uke 1. Eksepsjonsmeldingen oppgir hvor mye av uke 1 som er forfalt.",
    "7. Et forfalt planlagt ordremottak vises likevel i uke t, der behovet er, selv om det fysisk kommer tidligst i uke 1 + L. "
    "Forventet lagerbeholdning forutsetter altså at behovet dekkes i tide. Avviket varsles med eksepsjonen "
    "«Bestilling forsinket» eller «Bruk av sikkerhetslager».",
    "",
    "Eksepsjoner:",
    "• Bestilling forsinket: frigivelsen er forfalt (t − L < 1) og disponibel(t) < 0. Kunden får ikke levert i tide.",
    "• Bruk av sikkerhetslager: frigivelsen er forfalt (t − L < 1) og 0 ≤ disponibel(t) < SS. "
    "Lageret tærer på sikkerhetslageret før ny leveranse kan komme inn.",
    "• Mottak forsinket: nettobehov(t) > 0 og det finnes et planlagt mottak i en senere uke. "
    "Meldingen peker på første senere mottak, som bør fremskyndes.",
    "• Ledetid overskrider horisont: L ≥ H. Gjelder komponenten (uke «–»), uavhengig av behov.",
    "• En celle kan ha flere eksepsjoner. En komponent med minst én eksepsjon er kritisk.",
    "",
    "Status: verdiene er beregnet med skript og kontrollert for hånd (2026-10-19).",
]


def lag_fasit(produkt, mps, komponenter, mottak):
    wb = Workbook()
    h = len(mps)

    ws = wb.active
    ws.title = "Regler"
    ws["A1"] = f"Fasit: testeksempel 01 – {produkt}"
    ws["A1"].font = FONT_TITTEL
    r = tekstblokk(ws, 3, REGLER)
    for rr in range(3, r):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=8)
        ws.row_dimensions[rr].height = 28 if len(str(ws.cell(rr, 1).value or "")) > 110 else 15
    r += 1
    ws.cell(r, 1, "Hensikt med hver komponent i testeksempelet:").font = FONT_FET
    hode(ws, r + 1, ["Varenummer", "Beskrivelse", "Hensikt"])
    for i, k in enumerate(komponenter):
        celle(ws, r + 2 + i, 1, k[0])
        celle(ws, r + 2 + i, 2, k[1])
        celle(ws, r + 2 + i, 3, k[6])
        ws.merge_cells(start_row=r + 2 + i, start_column=3, end_row=r + 2 + i, end_column=8)
    bredder(ws, {c: 14 for c in "ABCDEFGH"})

    ws_t = wb.create_sheet("Fasittabell")
    hode(ws_t, 1, ["Varenummer", "Rad"] + [f"Uke {u}" for u in range(1, h + 1)])
    ws_e = wb.create_sheet("Eksepsjoner")
    hode(ws_e, 1, ["Varenummer", "Uke", "Eksepsjonstype", "Melding"])

    r, re_ = 2, 2
    for k in komponenter:
        rader, eks = beregn(mps, k, mottak)
        varsel_uker = {e[1] for e in eks}
        kritisk = bool(eks)
        for navn, verdier in rader.items():
            celle(ws_t, r, 1, k[0], FYLL_KRITISK if kritisk else None, fet=kritisk)
            celle(ws_t, r, 2, navn)
            for t, v in enumerate(verdier):
                fyll = FYLL_VARSEL if (t + 1) in varsel_uker and navn == "Nettobehov" else None
                celle(ws_t, r, 3 + t, v, fyll).alignment = Alignment(horizontal="right")
            r += 1
        r += 1  # tom rad mellom komponenter
        for e in eks:
            for i, v in enumerate(e):
                celle(ws_e, re_, i + 1, v).alignment = Alignment(wrap_text=True, vertical="top")
            re_ += 1
    bredder(ws_t, {"A": 14, "B": 26})
    for u in range(h):
        ws_t.column_dimensions[chr(ord("C") + u)].width = 8
    ws_t.freeze_panes = "C2"
    bredder(ws_e, {"A": 14, "B": 6, "C": 30, "D": 90})
    ws_e.freeze_panes = "A2"
    return wb


if __name__ == "__main__":
    lag_mal().save(UT / "COSI-MRP-mal.xlsx")
    if "--alle" in sys.argv:
        for k in KOMPONENTER:
            rader, eks = beregn(MPS, k, MOTTAK)
            print(k[0], k[6])
            for navn, v in rader.items():
                print(f"  {navn:28}", v)
            for e in eks:
                print("  !", e[1], e[2])
        # Ferdigproduktet uten lager og ledetid gir samme behov på nivå .1 som prognosen ganget med antall
        stykk = [("0", "100-001", PRODUKT, None, "Ferdigvare", 0, 0, 0)]
        stykk += [(".1", k[0], k[1], k[2], KJOPT, k[3], k[4], k[5]) for k in KOMPONENTER]
        lag_mal(MPS, stykk, MOTTAK).save(UT / "testeksempel-01.xlsx")
        lag_fasit(PRODUKT, MPS, KOMPONENTER, MOTTAK).save(UT / "testeksempel-01-fasit.xlsx")
